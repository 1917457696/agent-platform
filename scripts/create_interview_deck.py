from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor


OUT = Path(__file__).resolve().parents[1] / "output" / "huayi-world-tcm-ai-education-platform.pptx"
OUT.parent.mkdir(parents=True, exist_ok=True)

W, H = 13.333, 7.5
GREEN = "103A32"
GREEN_2 = "1E5A4A"
CREAM = "F7F1E4"
PAPER = "FBF8EF"
INK = "193630"
MUTED = "6D7B73"
RED = "C85D48"
GOLD = "D7C58F"
LINE = "D8D1C3"
WHITE = "FFFDF8"


def rgb(value):
    return RGBColor.from_string(value)


def box(slide, x, y, w, h, fill=None, line=None, radius=False, transparency=0):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                   Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill or PAPER)
    shape.fill.transparency = transparency
    shape.line.color.rgb = rgb(line or fill or PAPER)
    shape.line.width = Pt(0.8)
    return shape


def text(slide, value, x, y, w, h, size=16, color=INK, bold=False, font="Microsoft YaHei",
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, italic=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Pt(0)
    tf.margin_right = Pt(0)
    tf.margin_top = Pt(0)
    tf.margin_bottom = Pt(0)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)
    return shape


def line(slide, x1, y1, x2, y2, color=LINE, width=1.2, dash=None):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    if dash:
        shape.line.dash_style = dash
    return shape


def circle(slide, x, y, d, fill, label=None, label_color=WHITE, size=12):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(fill)
    if label:
        text(slide, label, x, y + d * .18, d, d * .6, size=size, color=label_color, bold=True, align=PP_ALIGN.CENTER)
    return shape


def base(slide, number, dark=False, kicker=None):
    bg = GREEN if dark else CREAM
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = rgb(bg)
    text(slide, "华医世界 · 中医课程 AI 产品方案", .55, .32, 5.5, .2, 8, GOLD if dark else RED, bold=True)
    text(slide, f"{number:02d}", 12.25, .32, .55, .2, 8, GOLD if dark else MUTED, bold=True, align=PP_ALIGN.RIGHT)
    if kicker:
        text(slide, kicker.upper(), .55, 6.98, 6, .2, 8, GOLD if dark else MUTED, bold=True)
    line(slide, .55, 6.83, 12.78, 6.83, "4C6C61" if dark else LINE, .7)


def title(slide, heading, sub=None, dark=False):
    text(slide, heading, .75, .95, 8.2, 1.05, 28, WHITE if dark else INK, bold=True, font="Microsoft YaHei")
    if sub:
        text(slide, sub, .77, 2.06, 8.7, .6, 12, "C7D6CC" if dark else MUTED)


def bullet(slide, label, body, x, y, w, color=INK, accent=RED):
    text(slide, label, x, y, w, .25, 12, accent, bold=True)
    text(slide, body, x, y + .32, w, .55, 11, color)


def add_slide(prs, dark=False, kicker=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    base(slide, len(prs.slides), dark=dark, kicker=kicker)
    return slide


prs = Presentation()
prs.slide_width = Inches(W); prs.slide_height = Inches(H)

# 1 cover
slide = add_slide(prs, dark=True, kicker="面试方案 · 2026.10")
text(slide, "中医课程\nAI 教学平台", .78, 1.18, 6.4, 1.75, 38, WHITE, bold=True)
text(slide, "从课程内容，到会讲解、会出题、会陪练、会反馈的学习服务", .82, 3.22, 6.7, .48, 15, "C7D6CC")
box(slide, 8.55, 1.18, 3.7, 3.8, fill=GREEN_2, line="356B5C", radius=True)
circle(slide, 9.15, 1.78, .58, RED, "课", size=14)
circle(slide, 10.9, 2.65, .58, GOLD, "练", label_color=GREEN, size=14)
circle(slide, 9.15, 3.54, .58, "6D9B86", "评", size=14)
line(slide, 9.73, 2.07, 10.9, 2.93, GOLD, 1.4)
line(slide, 10.9, 3.23, 9.73, 3.83, GOLD, 1.4)
line(slide, 9.15, 2.36, 9.15, 3.54, GOLD, 1.4)
text(slide, "课程讲解", 9.95, 1.88, 1.7, .24, 12, WHITE, bold=True)
text(slide, "医案陪练", 11.57, 2.76, 1.2, .24, 12, WHITE, bold=True)
text(slide, "学习反馈", 9.95, 3.64, 1.7, .24, 12, WHITE, bold=True)
text(slide, "问诊不是产品全部，而是临床思维训练模块", 8.84, 4.54, 3.0, .3, 10, GOLD, italic=True)

# 2 problem
slide = add_slide(prs, kicker="01 · 产品判断")
title(slide, "课程不是一段视频，\n也不是一个聊天框。", "中医教培的价值，应该从“看过”变成“学会、练会、能复盘”。")
items = [
    ("内容很多", "课程、医案、老师经验分散在视频、讲义和社群里，学员学完难以形成自己的知识结构。"),
    ("问题重复", "班主任和老师反复回答同一类问题，真正需要人工判断的高价值问题反而被淹没。"),
    ("学习断点", "学员看完课后缺少练习、反馈和下一步路径，不知道自己到底卡在哪里。"),
]
for i, (head, body) in enumerate(items):
    x = .78 + i * 4.1
    box(slide, x, 3.16, 3.75, 2.32, fill=PAPER, line=LINE)
    text(slide, f"0{i+1}", x+.22, 3.39, .4, .25, 11, RED, bold=True)
    text(slide, head, x+.22, 3.91, 3.1, .35, 19, INK, bold=True)
    text(slide, body, x+.22, 4.44, 3.15, .75, 11, MUTED)
text(slide, "产品机会", .8, 5.92, 1.2, .25, 11, RED, bold=True)
text(slide, "把一门课程升级为可互动、可练习、可追踪的 AI 学习服务。", 2.0, 5.9, 8.8, .32, 15, INK, bold=True)

# 3 product architecture
slide = add_slide(prs, dark=True, kicker="02 · 双产品架构")
title(slide, "一条课程主线，两个学习场景。", "课程 AI 教学平台是产品本体；医案/问诊陪练是其中一个高价值训练模块。", dark=True)
box(slide, .8, 2.95, 11.7, .78, fill=GREEN_2, line="477A69", radius=True)
text(slide, "课程内容层", 1.05, 3.16, 1.35, .25, 12, GOLD, bold=True)
text(slide, "讲义 · 名师医案 · 课程视频 · 训练营任务 · FAQ", 2.55, 3.16, 7.8, .25, 12, WHITE)
line(slide, 2.1, 3.73, 2.1, 4.4, GOLD, 1.2); line(slide, 6.65, 3.73, 6.65, 4.4, GOLD, 1.2); line(slide, 10.9, 3.73, 10.9, 4.4, GOLD, 1.2)
cols = [("AI 课程助教", "讲解、答疑、章节小测\n适用于所有学员", RED), ("医案陪练", "AI 扮演患者\n训练四诊与辨证思路", GOLD), ("教师协作", "审核反馈、看薄弱点\n沉淀 Bad Case", "6D9B86")]
for i, (head, body, col) in enumerate(cols):
    x = .8 + i * 4.25
    box(slide, x, 4.4, 3.72, 1.35, fill=col, line=col, radius=True)
    text(slide, head, x+.22, 4.68, 3.2, .25, 15, GREEN if col == GOLD else WHITE, bold=True)
    text(slide, body, x+.22, 5.04, 3.1, .5, 11, GREEN if col == GOLD else WHITE)
text(slide, "用户看到的是一套课程学习体验；系统内部按场景编排不同 Agent、检索和安全策略。", .82, 6.08, 10.8, .3, 12, "C7D6CC")

# 4 learning journey
slide = add_slide(prs, kicker="03 · 学习路径")
title(slide, "把“学完一节课”变成一条闭环。", "先帮助学员理解，再让学员主动练习，最后把反馈交给学员和老师。")
steps = [("01", "学懂", "按章节提问\nAI 用课程资料讲清楚"), ("02", "练习", "生成小测与复习卡\n检验是否真的理解"), ("03", "陪练", "进入医案\n主动问诊、组织证据"), ("04", "反馈", "指出遗漏与薄弱点\n给出下一步复习"), ("05", "协作", "教师审核与干预\nBad Case 反哺课程")]
for i, (num, head, body) in enumerate(steps):
    x = .78 + i * 2.42
    circle(slide, x+.58, 3.0, .58, RED if i == 2 else GREEN, num, size=10)
    if i < 4: line(slide, x+1.18, 3.29, x+2.37, 3.29, RED if i == 2 else LINE, 1.4)
    text(slide, head, x, 3.92, 2.1, .28, 17, INK, bold=True, align=PP_ALIGN.CENTER)
    text(slide, body, x+.04, 4.42, 2.05, .72, 11, MUTED, align=PP_ALIGN.CENTER)
box(slide, .82, 5.78, 11.6, .5, fill="F1E8D5", line="E4D7BD", radius=True)
text(slide, "核心变化：从“问 AI 得答案”，变成“学习行为被看见、被练习、被反馈”。", 1.1, 5.94, 11.0, .22, 13, GREEN, bold=True, align=PP_ALIGN.CENTER)

# 5 course assistant
slide = add_slide(prs, dark=True, kicker="04 · 产品一")
title(slide, "AI 课程助教：让老师的内容可以被反复学习。", "第一阶段面向一个课程主题，先验证“讲解 + 练习 + 来源”是否能降低重复答疑。", dark=True)
cards = [("课程问答", "围绕章节内容回答，\n每个答案带出处。"), ("换一种说法", "根据学员基础，\n把术语讲得更易懂。"), ("章节小测", "生成选择题、简答题\n与复习卡。"), ("复习路径", "根据错题和薄弱点，\n推荐下一章内容。")]
for i, (head, body) in enumerate(cards):
    x = .85 + (i % 2) * 5.9; y = 3.02 + (i // 2) * 1.42
    box(slide, x, y, 5.35, 1.08, fill=GREEN_2 if i == 0 else "174A40", line="4A7A6C", radius=True)
    text(slide, head, x+.25, y+.22, 1.55, .25, 14, GOLD if i == 0 else WHITE, bold=True)
    text(slide, body, x+2.0, y+.18, 2.95, .48, 11, "C7D6CC")
text(slide, "示例：学员问“什么是四诊合参？”→ 课程资料检索 → 通俗解释 → 章节出处 → 课后小测。", .86, 6.18, 11.2, .3, 12, GOLD)

# 6 clinical practice
slide = add_slide(prs, kicker="05 · 产品二")
title(slide, "医案陪练：把知识放进临床思维里。", "它不是在线问诊产品，而是课程平台里的一门“主动问、主动想、再复盘”的训练课。")
box(slide, .82, 3.0, 4.0, 2.8, fill=PAPER, line=LINE)
text(slide, "学员", 1.1, 3.3, .8, .3, 18, RED, bold=True)
text(slide, "“我想练习外感咳嗽医案，\n现在应该先问什么？”", 1.1, 3.9, 3.2, .75, 19, INK, bold=True)
text(slide, "主动提问 / 选择追问 / 提交判断", 1.1, 5.08, 3.2, .25, 11, MUTED)
box(slide, 8.5, 3.0, 4.0, 2.8, fill=GREEN, line=GREEN)
text(slide, "AI 患者 + 教练", 8.82, 3.3, 2.4, .3, 18, GOLD, bold=True)
text(slide, "只按病例状态逐步透露信息，\n不提前泄露标准答案。", 8.82, 3.9, 3.2, .75, 19, WHITE, bold=True)
text(slide, "追问 → 来源 → 反馈 → 教师审核", 8.82, 5.08, 3.2, .25, 11, "C7D6CC")
line(slide, 4.95, 4.4, 8.35, 4.4, RED, 2.2)
text(slide, "四诊信息", 5.65, 3.94, 1.5, .22, 10, RED, bold=True, align=PP_ALIGN.CENTER)
text(slide, "辨证依据", 6.15, 4.72, 1.5, .22, 10, RED, bold=True, align=PP_ALIGN.CENTER)
text(slide, "提交后才给反馈", .82, 6.16, 4, .25, 12, INK, bold=True)
text(slide, "遇到处方、剂量、个体化诊疗请求 → 安全闸 + 教师/医生审核", 6.0, 6.16, 6.0, .25, 11, RED, bold=True, align=PP_ALIGN.RIGHT)

# 7 difference
slide = add_slide(prs, dark=True, kicker="06 · 产品差异")
title(slide, "为什么不是通用大模型套壳？", "因为学习目标、知识来源、练习方式和安全边界都被产品化了。", dark=True)
diffs = [("有课程出处", "回答必须落到具体章节、讲义或老师医案，不只给一个看似合理的答案。"), ("有学习任务", "每次交互都服务于理解、练习或复习，不让聊天漫无目的地结束。"), ("有过程反馈", "评价学员问了什么、遗漏了什么、证据链是否完整，而不是只评价一句话。"), ("有教师入口", "高风险内容和 AI 草稿进入人工审核，教师可以看到问题分布并反哺课程。")]
for i, (head, body) in enumerate(diffs):
    x = .82 + (i % 2) * 6.0; y = 3.0 + (i // 2) * 1.62
    text(slide, f"0{i+1}", x, y, .5, .25, 11, GOLD, bold=True)
    text(slide, head, x+.62, y-.02, 2.9, .3, 16, WHITE, bold=True)
    text(slide, body, x+.62, y+.4, 4.75, .62, 11, "C7D6CC")
    line(slide, x+.62, y+1.23, x+5.6, y+1.23, "356B5C", .7)

# 8 tech
slide = add_slide(prs, kicker="07 · 技术落地")
title(slide, "AI 能力不是一个模型，而是一条可观察的工作流。", "我会把课程内容、业务规则和模型能力拆开，先做可运行 MVP，再按真实数据迭代。")
layers = [("内容层", "课程资料 / 医案 / FAQ / 题库", "E9DFC9"), ("检索层", "向量 + 关键词混合检索，保留出处", "DCE8D9"), ("编排层", "分流 → 检索 → 教练/答疑 → 安全审核", "D9E2E5"), ("产品层", "课程页 / 练习页 / 教师审核 / App API", "F0D8CF")]
for i, (head, body, col) in enumerate(layers):
    y = 2.85 + i * .76
    box(slide, 1.0, y, 11.1, .56, fill=col, line=col, radius=True)
    text(slide, head, 1.24, y+.17, 1.35, .22, 12, GREEN, bold=True)
    text(slide, body, 2.8, y+.17, 7.8, .22, 12, INK)
    text(slide, ["可审校", "可引用", "可追踪", "可接入"][i], 10.85, y+.17, .9, .22, 10, RED, bold=True, align=PP_ALIGN.RIGHT)
text(slide, "工程证据", .98, 6.12, 1.0, .22, 11, RED, bold=True)
text(slide, "React + FastAPI + Agent/RAG + SSE；本地 Trace 可展示节点、耗时、Token（精确或估算标注）；Embedding Key 通过会话级配置传入。", 2.1, 6.1, 10.2, .32, 11, MUTED)

# 9 MVP roadmap
slide = add_slide(prs, dark=True, kicker="08 · 0→1 计划")
title(slide, "先跑通一门课，再扩成平台。", "首期不做全量课程、完整 App 重构和自动诊疗，把边界留给真正的验证。", dark=True)
road = [("0—30 天", "找场景", "访谈讲师、班主任和学员\n选一个高频课程主题\n整理版权、权限和评测集"), ("31—60 天", "做 MVP", "课程问答 + 小测 + 一套医案\n来源引用、人工审核\n内部训练营试用"), ("61—90 天", "看数据", "小范围学员试用\n看答疑时间、采纳率、转人工\n再决定 App / 企微 / 小程序")]
for i, (time, head, body) in enumerate(road):
    x = .84 + i * 4.08
    box(slide, x, 3.0, 3.62, 2.45, fill=GREEN_2, line="467968", radius=True)
    text(slide, time, x+.25, 3.28, 1.5, .24, 11, GOLD, bold=True)
    text(slide, head, x+.25, 3.72, 2.4, .35, 20, WHITE, bold=True)
    text(slide, body, x+.25, 4.32, 3.0, .9, 11, "C7D6CC")
text(slide, "验证指标：重复问题减少量 · 班主任答疑时间 · 有来源回答比例 · 教师采纳率 · 转人工比例 · 练习完成率", .88, 6.03, 11.4, .28, 11, GOLD, bold=True)

# 10 demo evidence
slide = add_slide(prs, kicker="09 · 现场演示")
title(slide, "5 分钟演示：让业务看到“能跑、能解释、能迭代”。", "不临时开发完整 App，先用现有原型展示一条可信的课程学习链路。")
demo = [("1", "课程提问", "问：这一章的核心方法是什么？\n展示课程来源和通俗解释。"), ("2", "进入医案", "换季咳嗽合成医案。\nAI 先追问，学员可点选或补充。"), ("3", "提交判断", "学员说明依据。\n系统展示遗漏和临床思维反馈。"), ("4", "触发审核", "要求开方/剂量。\n安全闸转教师审核，不越界。")]
for i, (num, head, body) in enumerate(demo):
    x = .8 + i * 3.05
    circle(slide, x, 3.05, .5, RED if i == 3 else GREEN, num, size=11)
    text(slide, head, x, 3.82, 2.6, .3, 16, INK, bold=True)
    text(slide, body, x, 4.34, 2.45, .72, 11, MUTED)
    if i < 3: line(slide, x+.72, 3.3, x+2.8, 3.3, LINE, 1.3)
box(slide, .8, 5.76, 11.75, .56, fill="F1E8D5", line="E4D7BD", radius=True)
text(slide, "现场口径：当前使用合成资料和 Mock 外部系统，展示的是产品流程与工程能力，不冒充正式生产医疗系统。", 1.05, 5.95, 11.2, .22, 11, RED, bold=True, align=PP_ALIGN.CENTER)

# 11 role boundary
slide = add_slide(prs, dark=True, kicker="10 · 岗位边界")
title(slide, "我可以承担 0→1，但不会盲目承诺单人包办。", "先对结果负责，再把需要协作的边界讲清楚，反而更适合早期项目。", dark=True)
box(slide, .82, 3.0, 5.65, 2.65, fill=GREEN_2, line="477A69", radius=True)
text(slide, "我可以独立负责", 1.12, 3.35, 2.6, .3, 18, GOLD, bold=True)
text(slide, "场景调研与 MVP 定义\nAI 应用层：RAG、Agent、提示词、接口\n前端原型、流式体验、部署与 Trace\n评测集、Bad Case 和迭代文档", 1.12, 3.92, 4.7, 1.25, 12, WHITE)
box(slide, 6.83, 3.0, 5.65, 2.65, fill="174A40", line="477A69", radius=True)
text(slide, "正式上线需要协作", 7.13, 3.35, 2.8, .3, 18, GOLD, bold=True)
text(slide, "App 正式接入与账号体系\n数据治理、运维、成本与安全\n中医教师审核与课程版权\n规模化并发和产品运营", 7.13, 3.92, 4.7, 1.25, 12, "C7D6CC")
text(slide, "岗位沟通：如果职责是 AI 方向从 0 到 1，我期望转正 25K 左右；具体可结合首期目标、资源和绩效机制沟通。", .88, 6.1, 11.45, .3, 12, GOLD, bold=True)

# 12 close
slide = add_slide(prs, kicker="11 · 结论")
text(slide, "先做一门课，\n再做一套学习系统。", .82, 1.18, 7.3, 1.45, 34, INK, bold=True)
text(slide, "对华医世界，我建议首个产品从“课程 AI 助教 + 一套医案陪练”开始：\n范围可控、能快速上线、能被真实学员使用，也能为后续 App 产品化留下接口。", .84, 3.1, 6.7, .95, 14, MUTED)
box(slide, 8.35, 1.35, 3.9, 3.7, fill=GREEN, line=GREEN, radius=True)
text(slide, "课程资料", 8.76, 1.85, 2.9, .28, 17, WHITE, bold=True, align=PP_ALIGN.CENTER)
text(slide, "↓", 10.05, 2.3, .5, .3, 18, GOLD, bold=True, align=PP_ALIGN.CENTER)
text(slide, "AI 讲解", 8.76, 2.72, 2.9, .28, 17, WHITE, bold=True, align=PP_ALIGN.CENTER)
text(slide, "↓", 10.05, 3.17, .5, .3, 18, GOLD, bold=True, align=PP_ALIGN.CENTER)
text(slide, "医案练习", 8.76, 3.59, 2.9, .28, 17, WHITE, bold=True, align=PP_ALIGN.CENTER)
text(slide, "↓", 10.05, 4.04, .5, .3, 18, GOLD, bold=True, align=PP_ALIGN.CENTER)
text(slide, "反馈与教师审核", 8.76, 4.46, 2.9, .28, 17, GOLD, bold=True, align=PP_ALIGN.CENTER)
text(slide, "谢谢 / 欢迎现场演示", .85, 6.13, 4.1, .25, 12, RED, bold=True)

prs.save(OUT)
print(OUT)
